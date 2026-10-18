#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
validate_bid.py — 投标稿交付前校验（v2.0，双轨核查中的确定性轨）

用法：
    python validate_bid.py --bid ./bid --req requirements.csv
    python validate_bid.py --bid ./bid --req requirements.csv --map outline-map.csv
    python validate_bid.py --bid ./bid --req requirements.csv --score scoring-matrix.csv
    python validate_bid.py --bid ./bid --req requirements.csv \
        --project "XX系统集成项目" --project-no "KM-2026-0901" --report validate.json
    python validate_bid.py --bid ./bid --list-todo        # 只列待确认项

检查项（规则编号见 references/delivery-contract.md）：
    TODO-01   占位符残留
    COV-01    必答条目完全未出现
    COV-02    仅出现在覆盖声明行，正文无实质展开（v1 的漏洞：声明即算覆盖）
    COV-03    引用了台账中不存在的条目号
    GRAY-01   灰色响应表述（基本响应/部分响应…）
    NUM-01    章节编号重复/跳号/倒退（支持 第X章、# N.、N、三种体例）
    RISK-01   绝对化/宣传性表述
    RISK-02   注水套话
    MAP-01/02 映射表缺失必答条目 / chapter 列为空
    SCORE-01..06 评分点无章节、无证据、分值异常、把握度非法、拆分过粗、可达分偏低
    POLL-01..03  复用污染：旧项目编号、旧项目名称、内部用语
    FMT-01    篇幅统计（INFO；单篇过薄时提示）

    输入异常（路径不存在 / 无可读文件 / docx 依赖缺失）不产出 Finding，
    而是抛 InputError 并返回退出码 2 —— 因为此时"检查结论"本身无意义，
    继续输出 INFO/WARN 会造成"跑过了"的错觉。

退出码：0 = 通过；1 = 存在阻断项（或 --strict 下有 WARN）；2 = 输入错误
"""

import argparse
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

VERSION = "2.0.1"
TEXT_SUFFIXES = {".md", ".txt", ".markdown"}
SUPPORTED = TEXT_SUFFIXES | {".docx"}
SKIP_DIRS = {".git", "__pycache__", "node_modules", ".idea", ".vscode", ".venv"}

# ---------------------------------------------------------------- 规则表

TODO_PATTERNS = [
    (r"\[待确认", "待确认标记"),
    (r"\[待商务确认", "待商务确认标记"),
    (r"\[待决策人批准", "待决策批准标记"),
    (r"\[通识表述", "通识表述待技术确认"),
    (r"\bTODO\b", "TODO"),
    (r"\bTBD\b", "TBD"),
    (r"\bXXX\b", "XXX 占位"),
    (r"待补充", "待补充"),
    (r"待填写", "待填写"),
    (r"_{4,}", "下划线空位"),
    (r"[？?]{2,}", "连续问号占位"),
    (r"此处填写", "模板提示语"),
    (r"示例[:：]", "模板示例残留"),
]

GRAY_WORDS = [
    "基本响应", "基本满足", "部分响应", "原则上满足", "原则上响应",
    "大体满足", "视情况响应", "酌情满足", "基本符合", "勉强满足",
]

RISK_WORDS = [
    "100%", "零故障", "永不", "终身免费", "无限次", "彻底解决",
    "绝对安全", "万无一失", "卓越", "一流", "完美", "极致",
    "全方位", "深度融合", "业界领先", "国内领先",
]

DIFFUSION_WORDS = ["随着", "众所周知", "不言而喻", "势在必行", "日益凸显", "深入贯彻"]

INTERNAL_PHRASES = ["内部资料", "内部口径", "内部评审用", "待领导定", "按上次标书",
                    "上个项目", "沿用上次", "内部参考", "我司内部", "内部模板"]

NUM_UNIT_HINT = re.compile(r"\d+(?:\.\d+)?\s*(Gbps|Mbps|Tbps|路|台|个|套|年|天|小时|分钟|ms|%|TB|GB|PB)")

# 关键修正：\b 在中文语境下会失配（中文算 \w，导致"条目R-001"匹配不到）。
R_ID_RE = re.compile(r"(?<![A-Za-z0-9])R-\d{3,}(?![0-9])", re.IGNORECASE)
DECL_RE = re.compile(r"^\s*>?\s*覆盖条目\s*[：:](.*)$")
PROJ_NO_RE = re.compile(r"(?<![A-Za-z0-9])([A-Z]{2,8})[-_]?(\d{4})[-_]?(\d{2,6})(?![0-9])")
PROJ_NAME_RE = re.compile(r"([\u4e00-\u9fa5A-Za-z0-9（）()·]{2,24}项目)")
CN_NUM = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8,
          "九": 9, "十": 10, "零": 0, "百": 100}


# ---------------------------------------------------------------- 发现项

class Finding:
    __slots__ = ("level", "rule", "message", "file", "line", "evidence", "suggestion")

    def __init__(self, level, rule, message, file="", line=0, evidence="", suggestion=""):
        self.level = level
        self.rule = rule
        self.message = message
        self.file = file
        self.line = line
        self.evidence = evidence
        self.suggestion = suggestion

    def to_dict(self, fid):
        return {
            "id": f"F-{fid:03d}",
            "level": self.level,
            "rule": self.rule,
            "message": self.message,
            "location": {"file": self.file, "line": self.line},
            "evidence": self.evidence,
            "suggestion": self.suggestion,
        }


# ---------------------------------------------------------------- 读取

class InputError(Exception):
    pass


def read_any(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        try:
            import docx  # type: ignore
        except ImportError:
            raise InputError(
                f"读取 {path.name} 需要 python-docx，当前环境未安装。"
                "请 pip install python-docx，或把 docx 另存为 .md/.txt 后重试。"
                "（v1 会在此处静默生成'[跳过]'字符串并参与扫描，导致覆盖率误判）"
            )
        try:
            d = docx.Document(str(path))
        except Exception as e:  # noqa: BLE001
            raise InputError(f"无法解析 {path.name}：{e}（可能加密或损坏）")
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)

    for enc in ("utf-8-sig", "utf-8", "gb18030", "gbk"):
        try:
            return path.read_text(encoding=enc)
        except UnicodeDecodeError:
            continue
    return ""


def collect(bid_path: Path):
    files = {}
    if bid_path.is_file():
        files[bid_path.name] = read_any(bid_path)
        return files

    for p in sorted(bid_path.rglob("*")):
        if any(part in SKIP_DIRS for part in p.parts):
            continue
        if p.is_file() and p.suffix.lower() in SUPPORTED:
            files[str(p.relative_to(bid_path))] = read_any(p)
    return files


def load_req(req_path: Path):
    rows = None
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with req_path.open(encoding=enc, newline="") as f:
                rows = list(csv.DictReader(f))
            break
        except UnicodeDecodeError:
            continue
    if rows is None:
        raise InputError(f"无法读取 {req_path}")
    if not rows:
        raise InputError(f"{req_path} 为空或缺少表头")
    need = {"req_id", "mandatory"}
    missing = need - set(rows[0].keys())
    if missing:
        raise InputError(f"requirements.csv 缺少列：{sorted(missing)}")
    return rows


def load_csv_rows(path: Path):
    for enc in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open(encoding=enc, newline="") as f:
                return list(csv.DictReader(f))
        except UnicodeDecodeError:
            continue
    raise InputError(f"无法读取 {path}")


# ---------------------------------------------------------------- 文本扫描

def scan_patterns(files, patterns, level, rule, suggestion=""):
    hits = []
    compiled = [(re.compile(p, re.IGNORECASE), name) for p, name in patterns]
    for fname, text in files.items():
        for lineno, line in enumerate(text.splitlines(), 1):
            for pat, name in compiled:
                m = pat.search(line)
                if not m:
                    continue
                start = max(0, m.start() - 30)
                snippet = line[start:start + 90].strip()
                hits.append(Finding(level, rule, f"命中「{name}」", fname, lineno, snippet, suggestion))
    return hits


# ---------------------------------------------------------------- 检查项

def check_todo(files):
    return scan_patterns(files, TODO_PATTERNS, "BLOCKER", "TODO-01",
                         "交付前必须清零；无法确认的表述改写为不含承诺的写法，不要带标记交付")


def check_gray(files):
    return scan_patterns(files, [(w, w) for w in GRAY_WORDS], "BLOCKER", "GRAY-01",
                         "改为三值表达：完全响应并写明指标 / 负偏离并说明补偿 / 不响应")


def check_risk(files):
    out = scan_patterns(files, [(w, w) for w in RISK_WORDS], "WARN", "RISK-01",
                        "绝对化承诺须降级为可验证表述，或经决策人批准后保留")
    out += scan_patterns(files, [(w, w) for w in DIFFUSION_WORDS], "WARN", "RISK-02",
                         "删除背景铺陈，改为针对本项目的具体措施")
    return out


def check_coverage(files, req_rows):
    """双层覆盖：声明级（覆盖声明行）与实质级（正文其他位置）。"""
    declared, substantive, decl_loc = set(), set(), {}
    for fname, text in files.items():
        for lineno, line in enumerate(text.splitlines(), 1):
            ids = {m.upper() for m in R_ID_RE.findall(line)}
            if not ids:
                continue
            m = DECL_RE.match(line)
            if m:
                declared |= ids
                for i in ids:
                    decl_loc.setdefault(i, (fname, lineno))
            else:
                substantive |= ids

    all_ids = {r["req_id"].strip().upper() for r in req_rows if r.get("req_id")}
    mandatory = [r for r in req_rows
                 if (r.get("mandatory") or "").strip().upper() == "Y"
                 and not (r.get("dup_of") or "").strip()]

    findings = []
    for r in mandatory:
        rid = r["req_id"].strip().upper()
        if rid not in declared and rid not in substantive:
            findings.append(Finding(
                "BLOCKER", "COV-01",
                f"必答条目 {rid} 未在任何章节出现",
                evidence=(r.get("requirement") or "")[:90],
                suggestion="在 outline-map 指定章节补写响应段；若确认不响应，必须走偏离流程并说明理由",
            ))
        elif rid not in substantive:
            loc = decl_loc.get(rid, ("", 0))
            findings.append(Finding(
                "BLOCKER", "COV-02",
                f"必答条目 {rid} 仅出现在覆盖声明行，正文无实质展开",
                loc[0], loc[1], "> 覆盖条目：…（仅声明）",
                "在声明所指章节补写 ≥80 字响应段，必须含指标、措施、证据三者之一以上",
            ))

    extra = sorted((declared | substantive) - all_ids)
    for rid in extra:
        findings.append(Finding(
            "WARN", "COV-03", f"引用了台账中不存在的条目号 {rid}",
            suggestion="核对 req_id 拼写，或把该条补进 requirements.csv（可能是漏抽的必答项）",
        ))

    covered = len(mandatory) - len([f for f in findings if f.rule in ("COV-01", "COV-02")])
    pct = round(covered / len(mandatory) * 100, 1) if mandatory else 100.0
    stats = {
        "requirements_total": len(req_rows),
        "mandatory": len(mandatory),
        "coverage": {
            "declared": len(mandatory) - len([f for f in findings if f.rule == "COV-01"]),
            "substantive": covered,
            "percent": pct,
        },
    }
    return findings, stats


def _cn_to_int(s: str):
    """把「三」「十二」「二十三」「一百零五」这类章节号转成整数。"""
    if s.isdigit():
        return int(s)
    digits = {"零": 0, "一": 1, "二": 2, "三": 3, "四": 4, "五": 5,
              "六": 6, "七": 7, "八": 8, "九": 9}
    total, section, num = 0, 0, 0
    for ch in s:
        if ch in digits:
            num = digits[ch]
        elif ch == "十":
            section += (num or 1) * 10
            num = 0
        elif ch == "百":
            section += (num or 1) * 100
            num = 0
        else:
            return 0
    return total + section + num


def check_numbering(files):
    """一级章节编号连续性。优先 第X章，其次 '# N.'，再次 'N、'。"""
    findings = []
    for fname, text in files.items():
        cn, md, plain = [], [], []
        for lineno, line in enumerate(text.splitlines(), 1):
            s = line.strip()
            m = re.match(r"^#{0,6}\s*第([一二三四五六七八九十百零\d]+)[章篇部]", s)
            if m:
                cn.append((_cn_to_int(m.group(1)), lineno, s[:40]))
                continue
            m = re.match(r"^#\s+(\d+)\s*[\.、]", s)
            if m:
                md.append((int(m.group(1)), lineno, s[:40]))
                continue
            m = re.match(r"^(\d{1,2})、\s*\S", s)
            if m and len(s) <= 40:
                plain.append((int(m.group(1)), lineno, s[:40]))

        seq = cn if len(cn) >= 2 else (md if len(md) >= 2 else (plain if len(plain) >= 2 else []))
        if not seq:
            continue

        seen, prev = {}, 0
        for num, lineno, title in seq:
            if num in seen:
                findings.append(Finding("WARN", "NUM-01",
                                        f"一级章节编号重复：第 {num} 章（首次出现于第 {seen[num]} 行）",
                                        fname, lineno, title))
            elif prev and num > prev + 1:
                findings.append(Finding("WARN", "NUM-01",
                                        f"一级章节跳号：{prev} → {num}（缺 {num - prev - 1} 章）",
                                        fname, lineno, title))
            elif prev and num < prev:
                findings.append(Finding("WARN", "NUM-01",
                                        f"一级章节编号倒退：{prev} → {num}", fname, lineno, title))
            seen.setdefault(num, lineno)
            prev = max(prev, num)
    return findings


def check_map(map_path: Path, req_rows):
    findings = []
    try:
        mrows = load_csv_rows(map_path)
    except InputError as e:
        return [Finding("BLOCKER", "MAP-01", f"无法读取映射表：{e}")]

    mapped = {(r.get("req_id") or "").strip().upper() for r in mrows}
    for r in req_rows:
        if (r.get("mandatory") or "").strip().upper() != "Y":
            continue
        if (r.get("dup_of") or "").strip():
            continue
        rid = r["req_id"].strip().upper()
        if rid not in mapped:
            findings.append(Finding("BLOCKER", "MAP-01", f"必答条目 {rid} 未在 outline-map.csv 中映射章节",
                                    evidence=(r.get("requirement") or "")[:90],
                                    suggestion="G2 硬门禁：补映射后才能进入撰写"))
    for r in mrows:
        if not (r.get("chapter") or "").strip():
            findings.append(Finding("BLOCKER", "MAP-02", f"{r.get('req_id', '?')} 的 chapter 列为空",
                                    suggestion="填写章节号，否则等于没有落点"))
    return findings


def check_score(score_path: Path):
    findings = []
    try:
        rows = load_csv_rows(score_path)
    except InputError as e:
        return [Finding("BLOCKER", "SCORE-00", f"无法读取评分作战表：{e}")], {}

    theory, reachable, no_evidence = 0.0, 0.0, 0
    points = len(rows)
    no_rule = 0
    coef = {"高": 0.9, "中": 0.6, "低": 0.3, "无": 0.0}
    for r in rows:
        sid = (r.get("score_id") or "?").strip()
        item = (r.get("score_item") or "").strip()
        try:
            smax = float((r.get("score_max") or "0").strip() or 0)
        except ValueError:
            findings.append(Finding("WARN", "SCORE-03", f"{sid} 的 score_max 不是数字", evidence=str(r.get("score_max"))))
            smax = 0.0
        if smax <= 0:
            findings.append(Finding("WARN", "SCORE-03", f"{sid}（{item}）分值为 0 或非法", evidence=str(r.get("score_max"))))

        chap = (r.get("chapter") or "").strip()
        ev = (r.get("evidence") or "").strip()
        cap = (r.get("capability") or "").strip()
        rule = (r.get("score_rule") or "").strip()

        if not chap:
            findings.append(Finding("BLOCKER", "SCORE-01", f"{sid}（{item}）无承载章节",
                                    suggestion="在 outline-map 中指定章节并在该章展开，否则该分必丢"))
        if not ev:
            findings.append(Finding("BLOCKER", "SCORE-02", f"{sid}（{item}）证据材料为空",
                                    suggestion="填写 evidence；确无证据时写“无”并在待确认清单说明补齐路径"))
        elif ev == "无":
            no_evidence += 1
            findings.append(Finding("WARN", "SCORE-02", f"{sid}（{item}）标注为无证据，该项预计不得分",
                                    suggestion="确认是否接受丢分；否则补证据或改为正偏离争取部分分"))
        if cap and cap not in coef:
            findings.append(Finding("WARN", "SCORE-04", f"{sid} capability=「{cap}」不在 高/中/低/无 内"))
        if not rule:
            no_rule += 1
            findings.append(Finding("WARN", "SCORE-05", f"{sid}（{item}）score_rule 为空，无法核对是否漏拆",
                                    suggestion="补评分办法原文摘录；没有原文就无法判断这一条是否拆全"))

        theory += smax
        reachable += smax * coef.get(cap, 0.3)

    # 拆分过粗：一条评分点平均吃下 8 分以上，通常意味着把多个得分点压成了一行。
    # 用"平均分值"而不是单看条数——小项目评分项本来就少，单看条数会大面积误报。
    if points and points < 15 and theory / points > 8:
        findings.append(Finding("WARN", "SCORE-05",
                                f"评分点仅 {points} 条却摊到 {theory:g} 分（平均 {theory/points:.1f} 分/条），拆分过粗",
                                suggestion="按“评分办法原文一句话一条”重拆；一条吃 8 分以上通常说明有得分点被合并了"))

    if theory > 0:
        ratio = reachable / theory
        if ratio < 0.7:
            findings.append(Finding("INFO", "SCORE-06",
                                    f"可达分/理论满分 = {ratio:.0%}（<70%，竞争力不足）",
                                    suggestion="优先补 capability=低/无 且有证据可补的项"))
        stats = {"score_total": theory, "score_reachable": round(reachable, 1),
                 "score_ratio": round(ratio, 3), "score_no_evidence": no_evidence,
                 "score_points": points, "score_no_rule": no_rule}
    else:
        stats = {"score_points": points, "score_no_rule": no_rule} if points else {}
    return findings, stats


def check_pollution(files, project: str, project_no: str):
    findings = []
    pn = (project_no or "").strip().upper()
    core = (project or "").strip()
    for w in ("项目", "工程", "采购"):
        core = core.replace(w, "")

    for fname, text in files.items():
        for lineno, line in enumerate(text.splitlines(), 1):
            if pn:
                for m in PROJ_NO_RE.finditer(line):
                    hit = m.group(0).upper()
                    if hit.replace("-", "").replace("_", "") != pn.replace("-", "").replace("_", ""):
                        findings.append(Finding("BLOCKER", "POLL-01",
                                                f"出现非本项目编号 `{hit}`", fname, lineno, line.strip()[:90],
                                                "确认为复用残留，替换为本项目编号；若为引用他人案例，删除"))
            if core:
                for m in PROJ_NAME_RE.finditer(line):
                    name = m.group(1)
                    if any(k in name for k in ("本项目", "该项目", "本工程", "项目名称", "类似项目", "在建项目")):
                        continue
                    if core not in name:
                        findings.append(Finding("WARN", "POLL-02",
                                                f"出现疑似其他项目名称「{name}」", fname, lineno,
                                                line.strip()[:90], "核对是否复用残留；确为案例引用请注明“非本项目”"))

    findings += scan_patterns(files, [(w, w) for w in INTERNAL_PHRASES], "BLOCKER", "POLL-03",
                              "内部用语不得出现在投标文件中，改为对外表述")
    return findings


# ---------------------------------------------------------------- 输出

def hr(ch="─", n=68):
    return ch * n


def print_report(bid_path, req_path, files, findings, stats, failed, strict):
    blockers = [f for f in findings if f.level == "BLOCKER"]
    warns = [f for f in findings if f.level == "WARN"]
    infos = [f for f in findings if f.level == "INFO"]

    print(hr("="))
    print(f"投标文件交付前校验报告（v{VERSION}）")
    print(hr("="))
    print(f"稿件路径 : {bid_path.resolve()}")
    print(f"要求台账 : {Path(req_path).resolve()}（{stats.get('requirements_total', 0)} 条）")
    print(f"生成时间 : {dt.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    print(hr())

    print("【统计】")
    print(f"  · 稿件：{len(files)} 个文件 / {sum(len(t.splitlines()) for t in files.values())} 行 / "
          f"{sum(len(t) for t in files.values())} 字符")
    print(f"  · 必答条目：{stats.get('mandatory', 0)} 条")
    cov = stats.get("coverage", {})
    print(f"  · 必答覆盖（实质）：{cov.get('substantive', 0)}/{stats.get('mandatory', 0)}"
          f"（{cov.get('percent', 100.0)}%）")
    if "score_ratio" in stats:
        print(f"  · 评分可达分：{stats['score_reachable']}/{stats['score_total']}"
              f"（{stats['score_ratio']:.0%}），无证据项 {stats['score_no_evidence']} 个")
    num_count = len(NUM_UNIT_HINT.findall("\n".join(files.values())))
    print(f"  · 量化表述：{num_count} 处")
    print(hr())

    def dump(title, items, limit=30):
        if not items:
            return
        print(f"【{title}】共 {len(items)} 项")
        for f in items[:limit]:
            loc = f"{f.file}:{f.line}" if f.file else "—"
            print(f"  ▼ [{f.rule}] {f.message}")
            print(f"     位置 {loc}")
            if f.evidence:
                print(f"     证据 {f.evidence}")
            if f.suggestion:
                print(f"     整改 {f.suggestion}")
        if len(items) > limit:
            print(f"  … 另有 {len(items) - limit} 项未列出（见 --report 输出）")
        print(hr())

    dump("阻断项 BLOCKER", blockers)
    dump("提示 WARN", warns, 20)
    if infos:
        print("【信息 INFO】")
        for f in infos:
            print(f"  · [{f.rule}] {f.message}")
        print(hr())

    print(hr("="))
    if failed:
        print("结论：不通过 ✗")
        print("  先清零全部 BLOCKER，再重跑本脚本。")
        if warns:
            print("  WARN 不阻断，但需人工判断（绝对化承诺须决策人批准）。")
    else:
        print("结论：通过 ✓")
        print("  脚本只查形式。仍需按 compliance-checklist.md 做语义复核——")
        print("  尤其确认「答非所问」与「数值方向」，这两类脚本查不出。")
    print(hr("="))


def main():
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass

    ap = argparse.ArgumentParser(description="投标稿交付前合规校验（v2）")
    ap.add_argument("--bid", required=True, help="投标稿文件或目录")
    ap.add_argument("--req", help="requirements.csv 路径")
    ap.add_argument("--map", dest="map_file", help="outline-map.csv 路径（可选）")
    ap.add_argument("--score", dest="score_file", help="scoring-matrix.csv 路径（可选）")
    ap.add_argument("--project", help="本项目名称（用于复用污染比对）")
    ap.add_argument("--project-no", help="本项目编号（用于复用污染比对）")
    ap.add_argument("--report", help="输出机读结果 validate.json")
    ap.add_argument("--list-todo", action="store_true", help="只列出待确认项")
    ap.add_argument("--strict", action="store_true", help="把 WARN 也视为失败（等价 --fail-on warn）")
    ap.add_argument("--fail-on", choices=["blocker", "warn"], default="blocker",
                    help="CI 失败门限：blocker（默认）或 warn")
    ap.add_argument("--version", action="version", version=f"validate_bid.py {VERSION}")
    args = ap.parse_args()

    bid_path = Path(args.bid)
    if not bid_path.exists():
        print(f"[错误] 路径不存在：{bid_path}", file=sys.stderr)
        return 2

    try:
        files = collect(bid_path)
    except InputError as e:
        print(f"[错误] {e}", file=sys.stderr)
        return 2

    if not files:
        print(f"[错误] 在 {bid_path} 下未找到可读的 .md/.txt/.docx 文件", file=sys.stderr)
        return 2

    if args.list_todo:
        hits = check_todo(files)
        print(hr("="))
        print(f"待确认项清单（共 {len(hits)} 处）")
        print(hr("="))
        if not hits:
            print("✓ 未发现待确认标记")
        for f in hits:
            print(f"[{f.message}] {f.file}:{f.line}")
            print(f"    {f.evidence}")
        print(hr("="))
        return 0 if not hits else 1

    if not args.req:
        print("[错误] 需要 --req requirements.csv 才能做覆盖率检查", file=sys.stderr)
        return 2

    try:
        req_rows = load_req(Path(args.req))
    except InputError as e:
        print(f"[错误] {e}", file=sys.stderr)
        return 2

    findings = []
    findings += check_todo(files)
    cov_findings, stats = check_coverage(files, req_rows)
    findings += cov_findings
    findings += check_gray(files)
    findings += check_numbering(files)
    findings += check_risk(files)

    if args.map_file:
        findings += check_map(Path(args.map_file), req_rows)
    if args.score_file:
        score_findings, score_stats = check_score(Path(args.score_file))
        findings += score_findings
        stats.update(score_stats)
    if args.project or args.project_no:
        findings += check_pollution(files, args.project or "", args.project_no or "")

    # FMT-01 篇幅统计：永远只报 INFO。用途不是卡字数，而是发现"整篇只有标题"的稿子
    total_chars = sum(len(t) for t in files.values())
    avg = total_chars // max(len(files), 1)
    thin = sorted(((Path(n).name, len(t)) for n, t in files.items() if len(t) < 500),
                  key=lambda x: x[1])
    msg = f"正文 {len(files)} 篇 / {total_chars:,} 字，平均 {avg:,} 字每篇"
    if thin:
        msg += f"；{len(thin)} 篇不足 500 字：" + "、".join(f"{n}({c})" for n, c in thin[:5])
    findings.append(Finding(
        "INFO", "FMT-01", msg,
        suggestion="确认过薄的篇章是「故意精简」还是「只写了标题」——后者会让覆盖率虚高" if thin else "",
    ))

    blockers = [f for f in findings if f.level == "BLOCKER"]
    warns = [f for f in findings if f.level == "WARN"]
    failed = bool(blockers) or (args.strict and bool(warns)) or (args.fail_on == "warn" and bool(warns))

    print_report(bid_path, args.req, files, findings, stats, failed, args.strict)

    if args.report:
        out = Path(args.report)
        out.parent.mkdir(parents=True, exist_ok=True)
        payload = {
            "schema_version": "2.0",
            "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
            "tool": {"name": "validate_bid.py", "version": VERSION},
            "input": {
                "bid": str(bid_path),
                "requirements": str(args.req),
                "map": str(args.map_file or ""),
                "score": str(args.score_file or ""),
            },
            "stats": {
                "files": len(files),
                "lines": sum(len(t.splitlines()) for t in files.values()),
                "chars": sum(len(t) for t in files.values()),
                **stats,
            },
            "findings": [f.to_dict(i) for i, f in enumerate(findings, 1)],
            "summary": {
                "blocker": len(blockers),
                "warn": len(warns),
                "info": len([f for f in findings if f.level == "INFO"]),
                "conclusion": "不通过" if failed else "通过",
            },
        }
        out.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"[已写出] 机读结果：{out}")

    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
